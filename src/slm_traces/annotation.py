"""Blind annotation sheets (Excel with drop-downs) and label validation (plan §7).

The sheet hides the model name and the stratum; `trace_id` is an opaque hash, and the key
linking it back to model, item and partition is saved separately in data/annotation_set/.
"""

import random
import re
from collections import defaultdict
from pathlib import Path

import yaml

from . import ROOT

TAXONOMY_PATH = ROOT / "annotation" / "taxonomy.yaml"

TRACE_COLUMNS = [
    "trace_id", "task", "question", "options", "gold_answer", "model_final_answer", "answer_correct", "steps",
]
LABEL_COLUMNS = [
    "trace_valid", "first_error_step", "primary_failure_parent", "primary_failure_child",
    "propagated_error", "secondary_failure", "confidence", "notes",
]
STEP_LABEL_RE = re.compile(r"^\[(\d+)\] ", re.MULTILINE)
KEY_COLUMNS = ["trace_id", "dataset", "item_id", "model_id", "k", "partition"]


def load_taxonomy(path=TAXONOMY_PATH) -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def label_lists(taxonomy: dict) -> dict:
    parents = [p["id"] for p in taxonomy["parents"]]
    children = [c["id"] for p in taxonomy["parents"] for c in p["children"]]
    return {
        "trace_valid": ["TRUE", "FALSE"],
        "primary_failure_parent": parents,
        "primary_failure_child": children,
        "propagated_error": ["TRUE", "FALSE"],
        "secondary_failure": ["none"] + parents + children,
        "confidence": ["1", "2", "3"],
    }


def format_steps(steps: list[str]) -> str:
    return "\n".join(f"[{i}] {s}" for i, s in enumerate(steps, start=1))


def build_sheet(traces: list[dict], selection: list[dict], partition: str, seed: int):
    """Rows for one partition, grouped by question (questions and models in random order).

    Returns (sheet_rows, key_rows).
    """
    chosen = {(r["dataset"], r["item_id"]): r for r in selection if r["partition"] == partition}
    by_q = defaultdict(list)
    for t in traces:
        if (t["dataset"], t["item_id"]) in chosen:
            by_q[(t["dataset"], t["item_id"])].append(t)

    rng = random.Random(seed)
    questions = sorted(by_q)
    rng.shuffle(questions)
    sheet, key = [], []
    for q in questions:
        group = sorted(by_q[q], key=lambda t: t["trace_id"])
        rng.shuffle(group)
        for t in group:
            sheet.append({
                "trace_id": t["trace_id"],
                "task": t["task"],
                "question": t["question"],
                "options": "\n".join(f"{o['label']}. {o['text']}" for o in t.get("options") or []),
                "gold_answer": t["gold"],
                "model_final_answer": t["final_answer_text"],
                "answer_correct": "TRUE" if t["answer_correct"] else "FALSE",
                "steps": format_steps(t["steps"]),
                **{c: "" for c in LABEL_COLUMNS},
            })
            key.append({
                "trace_id": t["trace_id"], "dataset": t["dataset"], "item_id": t["item_id"],
                "model_id": t["model_id"], "k": chosen[q]["k"], "partition": partition,
            })
    return sheet, key


def export_xlsx(rows: list[dict], path, taxonomy: dict) -> None:
    """Write the sheet with drop-down lists for the label columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "traces"
    cols = TRACE_COLUMNS + LABEL_COLUMNS
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])

    # Drop-down values live on a hidden sheet (inline lists are limited to 255 characters).
    lists_ws = wb.create_sheet("lists")
    lists_ws.sheet_state = "hidden"
    n = max(len(rows), 1) + 1
    for j, (col, values) in enumerate(label_lists(taxonomy).items(), start=1):
        letter = get_column_letter(j)
        for i, v in enumerate(values, start=1):
            lists_ws.cell(row=i, column=j, value=v)
        dv = DataValidation(type="list", formula1=f"=lists!${letter}$1:${letter}${len(values)}", allow_blank=True)
        target = get_column_letter(cols.index(col) + 1)
        dv.add(f"{target}2:{target}{n}")
        ws.add_data_validation(dv)

    widths = {"question": 50, "options": 30, "steps": 80, "notes": 40, "model_final_answer": 18}
    label_fill = PatternFill("solid", fgColor="FFF2CC")
    for j, col in enumerate(cols, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = widths.get(col, 14)
        ws.cell(row=1, column=j).font = Font(bold=True)
        if col in LABEL_COLUMNS:
            ws.cell(row=1, column=j).fill = label_fill
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or str(v).strip() == ""


def _bool(v):
    s = str(v).strip().lower()
    return {"true": True, "false": False, "1": True, "0": False}.get(s)


def validate_row(row: dict, n_steps: int, taxonomy: dict) -> list[str]:
    """Check one annotated row against the codebook's consistency rules.

    first_error_step: -1 = no error; 1..n_steps = that step; n_steps + 1 = only the
    "Final answer:" line is wrong (e.g. it contradicts correct steps).
    """
    errs = []
    tid = row.get("trace_id", "?")
    parents = {p["id"]: {c["id"] for c in p["children"]} for p in taxonomy["parents"]}
    valid = _bool(row.get("trace_valid"))
    if valid is None:
        return [f"{tid}: trace_valid must be TRUE or FALSE"]
    try:
        step = int(float(row.get("first_error_step")))
    except (TypeError, ValueError):
        return [f"{tid}: first_error_step must be an integer"]
    parent, child = row.get("primary_failure_parent"), row.get("primary_failure_child")

    if valid:
        if step != -1:
            errs.append(f"{tid}: valid trace must have first_error_step = -1")
        if not _blank(parent) or not _blank(child):
            errs.append(f"{tid}: valid trace must not have a failure category")
        if _bool(row.get("answer_correct")) is False:
            errs.append(f"{tid}: answer is wrong, so the trace cannot be valid")
    else:
        if not 1 <= step <= n_steps + 1:
            errs.append(f"{tid}: first_error_step must be between 1 and {n_steps + 1}")
        if parent not in parents:
            errs.append(f"{tid}: primary_failure_parent '{parent}' is not in the taxonomy")
        elif not _blank(child) and child not in parents[parent]:
            errs.append(f"{tid}: child '{child}' does not belong to parent '{parent}'")
        if _bool(row.get("propagated_error")) is None:
            errs.append(f"{tid}: propagated_error must be TRUE or FALSE for an invalid trace")

    sec = row.get("secondary_failure")
    allowed_sec = {"none"} | set(parents) | {c for cs in parents.values() for c in cs}
    if not _blank(sec) and sec not in allowed_sec:
        errs.append(f"{tid}: secondary_failure '{sec}' is not in the taxonomy")
    conf = str(row.get("confidence")).strip().split(".")[0]
    if conf not in {"1", "2", "3"}:
        errs.append(f"{tid}: confidence must be 1, 2 or 3")
    elif conf == "1" and _blank(row.get("notes")):
        errs.append(f"{tid}: confidence 1 needs a note")
    return errs


def validate_sheet(rows: list[dict], taxonomy: dict) -> list[str]:
    """Validate every row that has been started; empty rows are reported as remaining."""
    errs, remaining = [], 0
    for row in rows:
        if all(_blank(row.get(c)) for c in LABEL_COLUMNS):
            remaining += 1
            continue
        n_steps = max((int(m) for m in STEP_LABEL_RE.findall(str(row.get("steps", "")))), default=0)
        errs += validate_row(row, n_steps, taxonomy)
    if remaining:
        errs.append(f"{remaining} trace(s) not annotated yet")
    return errs
